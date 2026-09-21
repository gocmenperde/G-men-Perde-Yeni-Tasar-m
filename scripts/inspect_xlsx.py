import openpyxl

hb_path = '/home/runner/workspace/attached_assets/HB_Duzeltilmis_Temiz_1781019586733.xlsx'
kit_path = '/home/runner/workspace/attached_assets/Kirtasiye-ve-Okul-Seti_2_1781019577099.xlsx'

print("=== HB Dosyasi ===")
wb_hb = openpyxl.load_workbook(hb_path, read_only=True, data_only=True)
print("Sheets:", wb_hb.sheetnames)
ws_hb = wb_hb.active
rows_hb = list(ws_hb.iter_rows(min_row=1, max_row=5, values_only=True))
for i, row in enumerate(rows_hb):
    print(f"Row {i}: {str(row)[:300]}")
print(f"Toplam satir: {ws_hb.max_row}, Toplam sutun: {ws_hb.max_column}")
wb_hb.close()

print("\n=== Kirtasiye Seti Dosyasi ===")
wb_kit = openpyxl.load_workbook(kit_path, read_only=True, data_only=True)
print("Sheets:", wb_kit.sheetnames)
ws_kit = wb_kit.active
rows_kit = list(ws_kit.iter_rows(min_row=1, max_row=5, values_only=True))
for i, row in enumerate(rows_kit):
    print(f"Row {i}: {str(row)[:300]}")
print(f"Toplam satir: {ws_kit.max_row}, Toplam sutun: {ws_kit.max_column}")
wb_kit.close()
